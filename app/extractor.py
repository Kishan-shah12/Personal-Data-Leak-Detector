import io
import re
from typing import Tuple


def extract_text_from_pdf(raw_bytes: bytes) -> Tuple[str, str, int]:
    """Extract text from PDF using pypdf."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw_bytes))
    page_texts = []
    total_pages = len(reader.pages)

    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            page_texts.append(f"--- [Page {idx}] ---\n{text}")

    full_text = "\n\n".join(page_texts)
    return full_text, f"PDF Document ({total_pages} page{'s' if total_pages != 1 else ''})", total_pages


def extract_text_from_docx(raw_bytes: bytes) -> Tuple[str, str, int]:
    """Extract text from Word .docx file using python-docx."""
    from docx import Document

    doc = Document(io.BytesIO(raw_bytes))
    sections = []

    # Extract paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            sections.append(p.text)

    # Extract tables
    table_count = len(doc.tables)
    for t_idx, table in enumerate(doc.tables, start=1):
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Remove duplicate adjacent cells from merged cells
                cleaned_cells = []
                for cell_val in row_cells:
                    if not cleaned_cells or cell_val != cleaned_cells[-1]:
                        cleaned_cells.append(cell_val)
                table_rows.append(" | ".join(cleaned_cells))
        if table_rows:
            sections.append(f"--- [Table {t_idx}] ---\n" + "\n".join(table_rows))

    full_text = "\n\n".join(sections)
    detail = f"Word Document (.docx) ({len(doc.paragraphs)} paragraphs, {table_count} tables)"
    return full_text, detail, len(sections)


def extract_text_from_xlsx(raw_bytes: bytes) -> Tuple[str, str, int]:
    """Extract text from Excel spreadsheet using openpyxl."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    sheet_texts = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows_text = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(val).strip() for val in row if val is not None and str(val).strip()]
            if cells:
                rows_text.append(", ".join(cells))
        if rows_text:
            sheet_texts.append(f"--- [Sheet: {sheet_name}] ---\n" + "\n".join(rows_text))

    wb.close()
    full_text = "\n\n".join(sheet_texts)
    detail = f"Excel Spreadsheet (.xlsx) ({len(wb.sheetnames)} sheet{'s' if len(wb.sheetnames) != 1 else ''})"
    return full_text, detail, len(sheet_texts)


def extract_text_from_rtf(raw_bytes: bytes) -> Tuple[str, str, int]:
    """Basic RTF text stripper."""
    text = raw_bytes.decode("latin-1", errors="replace")
    # Remove RTF control words
    clean = re.sub(r"\\[a-zA-Z0-9]+(?:-[0-9]+)? ?|\\[{}\\]", " ", text)
    clean = re.sub(r"[{}]", "", clean)
    clean = re.sub(r"\s+", " ", clean).strip()
    return clean, "Rich Text Format (.rtf)", 1


def extract_text_from_binary_fallback(raw_bytes: bytes) -> Tuple[str, str, int]:
    """Fallback strings extractor for binary documents."""
    # Find all printable strings of length >= 4
    strings = re.findall(rb"[\x20-\x7E]{4,}", raw_bytes)
    decoded_strings = [s.decode("ascii", errors="ignore") for s in strings]
    full_text = "\n".join(decoded_strings)
    return full_text, "Binary Document (Extracted Strings Stream)", len(decoded_strings)


def extract_document_text(filename: str, raw_bytes: bytes) -> Tuple[str, str, int]:
    """
    Extracts text from various document types based on file extension and signature.
    Returns: (extracted_text, document_type_description, sections_or_pages_count)
    """
    lower_name = (filename or "").lower().strip()
    ext = lower_name.rsplit(".", 1)[-1] if "." in lower_name else ""

    # Check for empty file
    if not raw_bytes:
        return "", "Empty File", 0

    # 1. PDF Documents
    if ext == "pdf" or raw_bytes.startswith(b"%PDF-"):
        try:
            return extract_text_from_pdf(raw_bytes)
        except Exception as e:
            # Fallback if pypdf encounters an issue
            text, _, count = extract_text_from_binary_fallback(raw_bytes)
            return text, f"PDF Document (Protected/Parsed via fallback: {str(e)[:40]})", count

    # 2. Word Documents (.docx)
    if ext in ("docx", "dotx", "docm"):
        try:
            return extract_text_from_docx(raw_bytes)
        except Exception as e:
            text, _, count = extract_text_from_binary_fallback(raw_bytes)
            return text, f"Word Document (.docx parsed via fallback: {str(e)[:40]})", count

    # 3. Excel Spreadsheets (.xlsx, .xlsm)
    if ext in ("xlsx", "xlsm", "xltx"):
        try:
            return extract_text_from_xlsx(raw_bytes)
        except Exception as e:
            text, _, count = extract_text_from_binary_fallback(raw_bytes)
            return text, f"Excel Spreadsheet (.xlsx parsed via fallback: {str(e)[:40]})", count

    # 4. RTF Documents
    if ext == "rtf" or raw_bytes.startswith(b"{\\rtf"):
        try:
            return extract_text_from_rtf(raw_bytes)
        except Exception:
            pass

    # 5. Legacy Word .doc or PowerPoint / binary formats
    if ext in ("doc", "ppt", "xls"):
        # Legacy binary OLE formats
        text, _, count = extract_text_from_binary_fallback(raw_bytes)
        return text, f"Legacy Office Document (.{ext} Extracted Text Stream)", count

    # 6. Plain text, source code, configs, logs, csv, json, md, html, etc.
    # Try utf-8 first
    try:
        decoded = raw_bytes.decode("utf-8")
        doc_type = f"Text Document ({ext.upper() if ext else 'Plain Text'})"
        return decoded, doc_type, 1
    except UnicodeDecodeError:
        pass

    # Try latin-1 / windows-1252
    try:
        decoded = raw_bytes.decode("latin-1")
        # Check if high proportion of non-printable chars (binary)
        printable_ratio = sum(c.isprintable() or c in "\r\n\t" for c in decoded[:2000]) / max(len(decoded[:2000]), 1)
        if printable_ratio < 0.7:
            # Likely binary
            return extract_text_from_binary_fallback(raw_bytes)
        return decoded, f"Text Document ({ext.upper() if ext else 'Encoded Text'})", 1
    except Exception:
        return extract_text_from_binary_fallback(raw_bytes)
